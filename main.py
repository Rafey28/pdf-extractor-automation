"""Batch processing entrypoint for PDF invoice extraction and Excel export."""
from __future__ import annotations

import logging
import time
import tempfile
import shutil
from pathlib import Path
from typing import List, Dict, Optional, Any
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pdf_extractor import PDFInvoiceExtractor, GeminiQuotaExceeded

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

from dotenv import load_dotenv
load_dotenv()


def ask_extraction_mode() -> Dict[str, Any]:
    """Ask user whether to use free regex model or AI-integrated Gemini model."""
    print("\n" + "="*60)
    print("PDF INVOICE EXTRACTOR - MODEL SELECTION")
    print("="*60)
    print("\n1. FREE MODEL (Regex-based)")
    print("   - No API key required")
    print("   - Fast extraction")
    print("   - Extracts: Invoice Date, Total Amount, Vendor Name")
    print("   - Fixed columns only")
    print("\n2. AI INTEGRATED MODEL (Google Gemini)")
    print("   - Requires GEMINI_API_KEY environment variable")
    print("   - Flexible custom columns")
    print("   - Custom extraction instructions")
    print("   - Professional formatting options")
    print("\n" + "="*60)

    while True:
        choice = input("\nSelect model (1 or 2): ").strip()
        if choice in ("1", "2"):
            break
        print("Invalid input. Please enter 1 or 2.")

    config: Dict[str, Any] = {"use_gemini": choice == "2"}

    if choice == "2":
        print("\n" + "="*60)
        print("GEMINI MODEL CONFIGURATION")
        print("="*60)

        # Ask for custom fields
        print("\nEnter custom columns to extract (comma-separated):")
        print("Examples: Invoice Date, Total Amount, Vendor Name, SKU Numbers, Purchase Order")
        fields_input = input("Columns: ").strip()
        custom_fields = [f.strip() for f in fields_input.split(",") if f.strip()]

        if not custom_fields:
            custom_fields = ["Invoice Date", "Total Amount", "Vendor Name"]
            print(f"Using default fields: {', '.join(custom_fields)}")

        config["custom_fields"] = custom_fields

        # Ask for extraction instructions
        print("\nEnter any specific extraction instructions (optional):")
        print("Leave blank to use default instructions.")
        instructions = input("Instructions: ").strip()
        if instructions:
            config["extraction_instructions"] = instructions

        # Ask for Excel formatting preferences
        print("\n" + "-"*60)
        print("EXCEL FILE FORMATTING OPTIONS")
        print("-"*60)
        print("\nWould you like to apply professional formatting?")
        format_choice = input("Apply formatting? (yes/no, default=yes): ").strip().lower()
        config["apply_formatting"] = format_choice != "no"

        if config["apply_formatting"]:
            print("\nFormatting will include:")
            print("  - Bold headers with light gray background")
            print("  - Auto-adjusted column widths")
            print("  - Center-aligned text")
            print("  - Currency formatting for amount fields")

    else:
        print("\nUsing FREE MODEL with fixed columns:")
        print("  - Filename, Date, Amount, Vendor")
        config["custom_fields"] = None

    return config


def process_pdfs(input_dir: Path, use_gemini: bool = False, custom_fields: Optional[List[str]] = None) -> pd.DataFrame:
    """Scan the input directory for PDF files, extract invoice data, and return a cleaned DataFrame.

    Implements rate limiting when using Gemini API to stay below free-tier RPM limits.
    """
    input_dir = input_dir.expanduser().resolve()
    pdf_files = sorted(input_dir.glob("*.pdf"))

    if not pdf_files:
        logger.warning("No PDF files found in %s", input_dir)
        if use_gemini and custom_fields:
            return pd.DataFrame(columns=["Filename"] + custom_fields)
        return pd.DataFrame(columns=["Filename", "Date", "Amount", "Vendor"])

    records: List[dict] = []
    total_files = len(pdf_files)

    for index, pdf_path in enumerate(pdf_files, start=1):
        logger.info("Processing %d/%d: %s", index, total_files, pdf_path.name)

        try:
            if use_gemini and custom_fields:
                extractor = PDFInvoiceExtractor(str(pdf_path), use_gemini=True)
                extracted = extractor.extract_data(custom_fields=custom_fields)
                record: Dict[str, Any] = {"Filename": pdf_path.name}
                record.update(extracted)  # type: ignore
            else:
                extractor = PDFInvoiceExtractor(str(pdf_path))
                extracted = extractor.extract_data()
                record = {
                    "Filename": pdf_path.name,
                    "Date": extracted.get("date"),
                    "Amount": extracted.get("amount"),
                    "Vendor": extracted.get("vendor"),
                }

            records.append(record)

            # Rate limiting: Add delay between Gemini API calls to stay below free-tier RPM limits
            if use_gemini and custom_fields and index < total_files:
                logger.debug("Rate limiting: waiting 4 seconds before next request")
                time.sleep(4)

        except GeminiQuotaExceeded:
            logger.error("Gemini API quota exceeded. Stopping PDF processing after %d/%d files.", index - 1, total_files)
            break

    df = pd.DataFrame(records)
    return normalize_dataframe(df, use_gemini=use_gemini, custom_fields=custom_fields)


def normalize_dataframe(df: pd.DataFrame, use_gemini: bool = False, custom_fields: Optional[List[str]] = None) -> pd.DataFrame:
    """Normalize DataFrame types and handle missing values gracefully."""
    if df.empty:
        return df

    df = df.copy()

    # For Gemini mode, do minimal normalization
    if use_gemini and custom_fields:
        # Just handle text fields
        for field in custom_fields:
            if field in df.columns and df[field].dtype == "object":
                df[field] = df[field].astype("string").str.strip().replace({"<NA>": ""})
        return df

    # For regex mode, use original normalization
    df["Amount"] = pd.to_numeric(df["Amount"], errors="coerce")
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Vendor"] = df["Vendor"].astype("string").str.strip().replace({"<NA>": ""})

    return df


def export_to_excel(df: pd.DataFrame, output_path: Path, apply_formatting: bool = True, max_retries: int = 3) -> None:
    """Export the DataFrame to Excel and apply professional formatting with openpyxl.

    Uses a temporary file to avoid conflicts with the file being open in Excel.
    Retries up to max_retries times if file operations fail.
    """
    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    last_error = None
    for attempt in range(max_retries):
        temp_path = None
        try:
            # Write to a temporary file first
            with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=output_path.parent, delete=False) as tmp:
                temp_path = Path(tmp.name)

            # Save initial DataFrame to Excel with pandas using openpyxl engine.
            df.to_excel(temp_path, index=False, engine="openpyxl")

            if apply_formatting:
                # Reopen the workbook with openpyxl to apply styling.
                workbook = load_workbook(temp_path)
                sheet = workbook.active
                assert sheet is not None

                header_font = Font(bold=True)
                header_fill = PatternFill(start_color="FFDDDDDD", end_color="FFDDDDDD", fill_type="solid")
                center_alignment = Alignment(horizontal="center", vertical="center")

                # Style the header row: bold and light gray fill.
                for cell in sheet[1]:
                    if cell.value is not None:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = center_alignment

                # Determine maximum width per column for auto-sizing.
                for column_cells in sheet.columns:
                    max_length = 0
                    column = column_cells[0].column_letter  # type: ignore
                    for cell in column_cells:
                        if cell.value is None:
                            continue
                        cell_value = str(cell.value)
                        max_length = max(max_length, len(cell_value))
                    adjusted_width = max_length + 2
                    sheet.column_dimensions[column].width = adjusted_width

                # Apply number formatting to Amount/Total columns if they exist.
                for cell in sheet[1]:
                    if cell.value is not None:
                        cell_name = str(cell.value).strip().lower()
                        if "amount" in cell_name or "total" in cell_name or "price" in cell_name:
                            amount_column = cell.column_letter
                            for row in range(2, sheet.max_row + 1):
                                amount_cell = sheet[f"{amount_column}{row}"]
                                if amount_cell.value is not None:
                                    try:
                                        amount_cell.number_format = "$#,##0.00"
                                    except Exception:
                                        pass

                # Center-align Date columns.
                for cell in sheet[1]:
                    if cell.value is not None:
                        cell_name = str(cell.value).strip().lower()
                        if "date" in cell_name:
                            date_column = cell.column_letter
                            for row in range(2, sheet.max_row + 1):
                                date_cell = sheet[f"{date_column}{row}"]
                                date_cell.alignment = center_alignment

                workbook.save(temp_path)
                workbook.close()

            # Use shutil.move which handles Windows file replacement
            shutil.move(str(temp_path), str(output_path))

            logger.info("Exported styled Excel report to %s", output_path)
            return

        except PermissionError as e:
            last_error = e
            # Clean up temp file if it exists
            if temp_path and temp_path.exists():
                try:
                    temp_path.unlink()
                except Exception:
                    pass

            if attempt < max_retries - 1:
                wait_time = 2 ** attempt
                logger.warning(
                    "File locked (attempt %d/%d). Waiting %d seconds...",
                    attempt + 1, max_retries, wait_time
                )
                time.sleep(wait_time)
            else:
                logger.error(
                    "Failed to write Excel file after %d attempts. "
                    "Please close the file if it's open in Excel.",
                    max_retries
                )
        except Exception as e:
            # Clean up temp file on unexpected error
            if temp_path and temp_path.exists():
                try:
                    temp_path.unlink()
                except Exception:
                    pass
            logger.error("Unexpected error while exporting to Excel: %s", e)
            raise

    if last_error:
        raise PermissionError(
            f"Could not write to {output_path}. The file may be open in Excel. "
            f"Close it and try again. Error: {last_error}"
        )


def main() -> None:
    """Main entry point with user configuration."""
    # Get user's model choice
    config = ask_extraction_mode()

    input_dir = Path("./data/input_pdfs")
    output_path = Path("./data/output/extracted_report.xlsx")

    logger.info("Starting batch PDF extraction")
    logger.info("Mode: %s", "Gemini AI" if config["use_gemini"] else "Free Regex")

    try:
        # Process PDFs based on selected mode
        df = process_pdfs(
            input_dir,
            use_gemini=config["use_gemini"],
            custom_fields=config.get("custom_fields")
        )
    except GeminiQuotaExceeded as e:
        print("\n" + "="*60)
        print("ERROR: GEMINI API QUOTA EXCEEDED")
        print("="*60)
        print(f"Error: {e}")
        print("\nPDF processing has been stopped to avoid further API errors.")
        print("Please wait before attempting again or check your Gemini API quota.")
        print("="*60 + "\n")
        logger.error("Batch processing stopped due to Gemini quota: %s", e)
        return

    if df.empty:
        logger.info("No extracted records to export. Exiting.")
        return

    # Export with formatting option
    apply_formatting = config.get("apply_formatting", True)
    export_to_excel(df, output_path, apply_formatting=apply_formatting)

    print("\n" + "="*60)
    print("✓ BATCH PROCESSING COMPLETED SUCCESSFULLY")
    print("="*60)
    print(f"Output file: {output_path}")
    print(f"Records extracted: {len(df)}")
    print(f"Columns: {', '.join(df.columns)}")
    print("="*60 + "\n")

    logger.info("Batch processing completed successfully.")


if __name__ == "__main__":
    main()
